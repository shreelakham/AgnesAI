"""Generate a pitch-deck .pptx from a pitch dict (and an optional concept image).

Design: deep-teal/navy palette, dark title + close ("sandwich"), light content
slides, large stat callouts for impact, no AI-tell accent stripes/underlines.
Pure python-pptx so the whole backend stays in Python.
"""

from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- Palette --------------------------------------------------------------
NAVY = RGBColor(0x1B, 0x2A, 0x41)
TEAL = RGBColor(0x02, 0x80, 0x90)
MINT = RGBColor(0x02, 0xC3, 0x9A)
INK = RGBColor(0x1A, 0x2B, 0x32)
MUTED = RGBColor(0x6B, 0x7C, 0x85)
LIGHT = RGBColor(0xF4, 0xF7, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xFF, 0xFF, 0xFF)

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

SEV_COLORS = {"high": RGBColor(0xD7, 0x4C, 0x4C),
              "medium": RGBColor(0xE0, 0x9B, 0x2D),
              "low": RGBColor(0x4C, 0x9A, 0x6B)}


def _clip(text: str, n: int) -> str:
    text = (text or "").strip()
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _text(slide, x, y, w, h, text, size, color, *, bold=False, font=BODY_FONT,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, line=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(0)
    tf.margin_top = tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    return box


def _card(slide, x, y, w, h, fill=CARD):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = RGBColor(0xE2, 0xE8, 0xEB)
    shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def _circle_label(slide, x, y, d, text, fill=TEAL, txt=WHITE):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                               Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.name = BODY_FONT
    r.font.color.rgb = txt
    return c


# ---------------------------------------------------------------------------
def build_deck(pitch: dict, image_bytes: bytes | None = None) -> bytes:
    """Return the .pptx as bytes."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    W = 13.333

    brand = pitch.get("brand", "the brand")

    # ---- Slide 1: Title (dark) -------------------------------------------
    s = prs.slides.add_slide(blank)
    _bg(s, NAVY)
    img_w = 5.2
    _text(s, 0.7, 1.0, 6.6, 0.5, "ANYMIND  ·  COMPETITIVE PITCH",
          14, MINT, bold=True, font=BODY_FONT)
    _text(s, 0.7, 1.7, 6.8, 1.6, brand, 46, WHITE, bold=True, font=HEAD_FONT)
    _text(s, 0.7, 3.5, 6.6, 2.4,
          _clip(pitch.get("headline_pitch", ""), 220),
          20, RGBColor(0xCA, 0xDC, 0xFC), italic=True, line=1.15)
    if image_bytes:
        s.shapes.add_picture(BytesIO(image_bytes),
                             Inches(W - img_w - 0.6), Inches(1.6),
                             height=Inches(4.3))
    else:
        _card(s, W - img_w - 0.6, 1.6, img_w, 4.3, fill=TEAL)

    # ---- Slide 2: Landscape (light) --------------------------------------
    s = prs.slides.add_slide(blank)
    _bg(s, LIGHT)
    _text(s, 0.7, 0.55, 11, 0.8, "The Landscape", 38, INK, bold=True, font=HEAD_FONT)
    _text(s, 0.7, 1.7, 7.0, 0.4, "WHO THEY ARE", 14, TEAL, bold=True)
    _text(s, 0.7, 2.15, 7.0, 1.7, _clip(pitch.get("brand_summary", ""), 380),
          15, INK, line=1.2)
    _text(s, 0.7, 4.05, 7.0, 0.4, "COMPETITIVE POSITION", 14, TEAL, bold=True)
    _text(s, 0.7, 4.5, 7.0, 2.2, _clip(pitch.get("competitive_position", ""), 380),
          15, INK, line=1.2)
    # competitor chips
    _text(s, 8.2, 1.7, 4.4, 0.4, "KEY COMPETITORS", 14, TEAL, bold=True)
    cy = 2.2
    for comp in pitch.get("likely_competitors", [])[:5]:
        _card(s, 8.2, cy, 4.4, 0.62)
        _text(s, 8.45, cy + 0.13, 4.0, 0.4, _clip(comp, 48), 14, INK, bold=True,
              anchor=MSO_ANCHOR.MIDDLE)
        cy += 0.78

    # ---- Slide 3: Gaps (light) -------------------------------------------
    s = prs.slides.add_slide(blank)
    _bg(s, LIGHT)
    _text(s, 0.7, 0.55, 11, 0.8, "Where They're Losing Ground", 36, INK,
          bold=True, font=HEAD_FONT)
    gaps = pitch.get("gaps", [])[:3]
    gw = 3.86
    gx = 0.7
    for g in gaps:
        _card(s, gx, 1.7, gw, 4.6)
        sev = (g.get("severity") or "medium").lower()
        badge = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(gx + 0.3), Inches(2.0),
                                   Inches(1.5), Inches(0.42))
        badge.fill.solid()
        badge.fill.fore_color.rgb = SEV_COLORS.get(sev, MUTED)
        badge.line.fill.background()
        badge.shadow.inherit = False
        bp = badge.text_frame.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        br = bp.add_run(); br.text = sev.upper()
        br.font.size = Pt(11); br.font.bold = True
        br.font.color.rgb = WHITE; br.font.name = BODY_FONT
        _text(s, gx + 0.3, 2.65, gw - 0.6, 1.4, _clip(g.get("gap", ""), 90),
              18, INK, bold=True, font=HEAD_FONT, line=1.1)
        _text(s, gx + 0.3, 4.1, gw - 0.6, 2.0, _clip(g.get("why_it_matters", ""), 240),
              13, MUTED, line=1.18)
        gx += gw + 0.18

    # ---- Slide 4: Solution mapping (light) -------------------------------
    s = prs.slides.add_slide(blank)
    _bg(s, LIGHT)
    _text(s, 0.7, 0.55, 11, 0.8, "How AnyMind Closes the Gaps", 36, INK,
          bold=True, font=HEAD_FONT)
    sy = 1.7
    for m in pitch.get("solution_mapping", [])[:4]:
        rh = 1.18
        _card(s, 0.7, sy, 11.9, rh)
        # product pill (handles variable-length names cleanly)
        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(0.95), Inches(sy + 0.34),
                                  Inches(1.55), Inches(0.5))
        pill.fill.solid(); pill.fill.fore_color.rgb = TEAL
        pill.line.fill.background(); pill.shadow.inherit = False
        pp = pill.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
        pr = pp.add_run(); pr.text = _clip(m.get("product", ""), 14)
        pr.font.size = Pt(13); pr.font.bold = True
        pr.font.color.rgb = WHITE; pr.font.name = BODY_FONT
        _text(s, 2.75, sy + 0.16, 2.7, 0.9,
              _clip(m.get("gap_addressed", ""), 60), 13, TEAL, bold=True,
              anchor=MSO_ANCHOR.MIDDLE, line=1.05)
        _text(s, 5.7, sy + 0.16, 6.7, 0.9,
              _clip(m.get("how_it_helps", ""), 190), 13, INK,
              anchor=MSO_ANCHOR.MIDDLE, line=1.1)
        sy += rh + 0.16

    # ---- Slide 5: Projected impact (light, big stats) --------------------
    s = prs.slides.add_slide(blank)
    _bg(s, LIGHT)
    _text(s, 0.7, 0.55, 11, 0.8, "Projected Impact", 38, INK, bold=True, font=HEAD_FONT)
    imp = pitch.get("projected_impact", {})
    rev = imp.get("revenue_or_profit_uplift", {})
    eff = imp.get("time_or_efficiency_gain", {})

    def stat(x, big, label, sub):
        _card(s, x, 1.8, 3.86, 4.3)
        # auto-size the headline number so long ranges don't overflow
        n = len(big)
        size = 50 if n <= 7 else (38 if n <= 12 else 28)
        _text(s, x + 0.2, 2.25, 3.46, 1.5, big, size, TEAL, bold=True,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=HEAD_FONT)
        _text(s, x + 0.3, 4.0, 3.26, 0.5, label, 16, INK, bold=True,
              align=PP_ALIGN.CENTER)
        _text(s, x + 0.3, 4.55, 3.26, 1.4, sub, 12, MUTED,
              align=PP_ALIGN.CENTER, line=1.2)

    stat(0.7, f"{rev.get('low','—')}–{rev.get('high','—')}",
         "Revenue / profit uplift",
         f"{_clip(rev.get('basis',''),150)}  ({rev.get('timeframe','')})")
    stat(4.74, f"{eff.get('low','—')}–{eff.get('high','—')}",
         "Time / efficiency gain", _clip(eff.get("what_it_speeds_up", ""), 150))
    stat(8.78, _clip(imp.get("time_to_value", "—"), 14),
         "Time to value", "From kickoff to measurable results.")

    # ---- Slide 6: The package / close (dark) -----------------------------
    s = prs.slides.add_slide(blank)
    _bg(s, NAVY)
    pkg = pitch.get("recommended_package", {})
    _text(s, 0.7, 0.9, 11, 0.5, "THE RECOMMENDED PACKAGE", 14, MINT, bold=True)
    _text(s, 0.7, 1.5, 11.9, 1.1, _clip(pkg.get("package_name", "Growth Package"), 60),
          40, WHITE, bold=True, font=HEAD_FONT)
    _text(s, 0.7, 2.9, 11.9, 1.6, _clip(pkg.get("rationale", ""), 320),
          17, RGBColor(0xCA, 0xDC, 0xFC), line=1.2)
    # product chips
    px = 0.7
    for prod in pkg.get("included_products", [])[:6]:
        chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(px), Inches(4.7), Inches(1.9), Inches(0.6))
        chip.fill.solid(); chip.fill.fore_color.rgb = TEAL
        chip.line.fill.background(); chip.shadow.inherit = False
        cp = chip.text_frame.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run(); cr.text = _clip(prod, 16)
        cr.font.size = Pt(14); cr.font.bold = True
        cr.font.color.rgb = WHITE; cr.font.name = BODY_FONT
        px += 2.05
    _text(s, 0.7, 5.7, 11.9, 1.2,
          _clip(pkg.get("suggested_engagement", ""), 220), 14,
          RGBColor(0x9F, 0xB4, 0xC4), italic=True, line=1.2)

    out = BytesIO()
    prs.save(out)
    return out.getvalue()